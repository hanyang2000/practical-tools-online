"""天猫首页数据分析 - 四表解析器。

把《2026天猫旗舰店页面数据监测.xlsx》解析成统一结构，供字段映射预览与后续面板消费。

对应用户实际导出结构：
- 整体数据：两行合并表头（人群×指标），A=日期（Excel 序列或日期），B=活动名（向下延续），
  季度/H1/H2 处插入「日常平均数据 / 活动平均数据 / 汇总」行。
- 日常页面数据 / 活动页面数据：按「活动块」组织，每块 = 标题行 + 3 行表头 + 数据行；
  A=板块（向下延续）、B=子模块、C/D=平均数据[点击率,支付转化]、
  其后每日期 4 列 = 新客[点击率,支付转化] / 老客[点击率,支付转化]。
- 月度数据对比：样例已移除，v1 不消费。

输出：parse_file() 返回 dict：
  {"sheets": {表名: {kind, fields, data, sample, anomalies}}, "unrecognized_sheets": [...]}
"""
import datetime
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

# ---------------- 通用 ----------------

_DASH = re.compile(r'^[\s\-/—#\\]+$')


def _clean(v):
    """单元格值归一化：空值/'-'/'/' 等 → None；数值 → int/float；日期 → ISO 字符串；其余 → 字符串。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return v if v == v else None  # NaN → None
    s = str(v).strip()
    if not s or _DASH.match(s) or s.upper() in ("#REF!", "#N/A", "#VALUE!", "NULL"):
        return None
    try:
        return float(s)
    except ValueError:
        return s


def _date_to_iso(v):
    """把日期单元格转成 'YYYY-MM-DD'：兼容 datetime / date / Excel 序列 / '2026/1/7 官宣前' 这类字符串。"""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, str):
        m = re.search(r'(\d{4})[-/](\d{1,2})[-/](\d{1,2})', v)
        if m:
            y, mo, d = (int(x) for x in m.groups())
            try:
                return datetime.date(y, mo, d).isoformat()
            except ValueError:
                return None
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < v < 80000:
        try:
            return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))).isoformat()
        except Exception:
            return None
    return None


# ---------------- 整体数据 ----------------

SUMMARY_B_TYPES = {
    "日常平均数据": "日常平均",
    "活动平均数据": "活动平均",
    "汇总（区分新老客）": "汇总(区分新老客)",
    "汇总（不区分新老客）": "汇总(不区分新老客)",
}
_GROUP_MARKERS = ("新客", "老客")


def parse_overall(ws):
    max_col = ws.max_column
    # 1) 人群分组的列边界：从第 1 行合并表头推导（如 D1:T1=新客、U1:AK1=老客）
    group_start = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row == 1 and rng.min_col > 1:
            v = str(ws.cell(1, rng.min_col).value or "").strip()
            if v in _GROUP_MARKERS:
                group_start.setdefault(v, rng.min_col)
    old_start = group_start.get("老客")
    if not old_start:
        old_start = max_col + 1
    # 2) 列映射：第 2 行指标名 + 人群归属（「总访客数」独立，不属于新客/老客任何人群）
    col_map = []  # (col, group, metric)；总访客数列 group 标为「整体」
    for c in range(3, max_col + 1):
        m = str(ws.cell(2, c).value or "").strip()
        if not m:
            continue
        group = "老客" if c >= old_start else "新客"
        if m == "总访客数":
            group = "整体"
        col_map.append((c, group, m))

    anomalies = []

    # 3) 逐行解析
    rows = []
    activity = None
    for r in range(3, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        b_s = str(b).strip() if b else ""
        iso = _date_to_iso(a)
        if iso:  # 日数据行
            if b_s:
                activity = b_s
            seg = {"新客": {}, "老客": {}, "整体": {}}
            for c, g, m in col_map:
                v = _clean(ws.cell(r, c).value)
                if v is None:
                    continue
                if g == "整体":
                    seg["整体"][m] = v
                else:
                    seg[g][m] = v
            # 总访客数 = 新客访客数 + 老客访客数（源表可能没有该列，程序自行计算）
            total = seg["整体"].get("总访客数")
            nk_visit = seg["新客"].get("访客数")
            ok_visit = seg["老客"].get("访客数")
            if nk_visit is not None and ok_visit is not None:
                total = nk_visit + ok_visit
            seg["整体"]["总访客数"] = total
            rows.append({"date": iso, "activity": activity, "is_summary": False,
                         "period_label": None, "summary_type": None, "segments": seg})
        else:  # 汇总/平均行
            st = SUMMARY_B_TYPES.get(b_s)
            if st:
                rows.append({"date": None, "activity": None, "is_summary": True,
                             "period_label": str(a).strip() if a else None,
                             "summary_type": st, "segments": {"新客": {}, "老客": {}, "整体": {}}})
            # 其它行（空行/尾部）忽略

    # 4) 统计
    daily = [x for x in rows if not x["is_summary"]]
    summary = [x for x in rows if x["is_summary"]]
    years = sorted({x["date"][:4] for x in daily if x.get("date")})
    activities = []
    for x in daily:
        if x["activity"] and x["activity"] not in activities:
            activities.append(x["activity"])
    fields = {
        "segments": ["新客", "老客", "整体"],
        "columns": [{"col": get_column_letter(c), "group": g, "metric": m} for c, g, m in col_map],
        "activities": activities,
        "daily_count": len(daily),
        "summary_count": len(summary),
        "summary_types": [s["summary_type"] for s in summary],
        "date_range": {"min": daily[0]["date"], "max": daily[-1]["date"]} if daily else None,
        "years": years,
        "note_total_visitors": "总访客数独立于人群，程序按 新客访客数+老客访客数 计算（源表无此列也能算）",
    }
    return {"kind": "overall", "fields": fields, "data": rows,
            "sample": rows[:5], "anomalies": anomalies}


# ---------------- 日常 / 活动页面数据 ----------------

# 波次/形态关键词：识别「同一活动的不同波段」时剔除，归并到基础活动名
WAVE_KEYWORDS = [
    "官宣预热期", "官宣前", "官宣后", "官宣", "最后一波", "第一波", "第二波",
    "第三波", "第四波", "倒计时", "返场", "抢先购", "焕新周", "预热",
    "预售", "开场", "爆发", "尾款", "余热", "新品", "上新前", "上新", "预热期",
]


def _normalize_activity_name(name: str) -> str:
    """把块名归一化成基础活动名：去括号等级、去波次关键词、去「活动-」分类前缀。"""
    s = name.strip()
    s = re.sub(r'[（(][^）)]*[）)]', '', s)          # 去括号等级（S/S-/常规/自制S-…）
    for kw in WAVE_KEYWORDS:
        s = s.replace(kw, '')
    s = re.sub(r'^活动[\s\-—]*', '', s)              # 2025 文件的「活动-」分类前缀
    s = re.sub(r'[\s\-—]+', ' ', s).strip()           # 分隔符归一
    return s or name.strip()


def _activity_grade(name: str) -> str:
    """从活动名取等级：括号内容（S+/S/S-/常规/自制S-…）；无括号返回空串。

    等级只用于展示/筛选标注，不进归组键（`_activity_group_key` 会去括号）。
    """
    if not name:
        return ""
    m = re.search(r'[（(]([^）)]*)[）)]', name)
    return m.group(1).strip() if m else ""


def _block_grade(title: str, name: str = "") -> str:
    """从块标题/名提取活动等级，兼容 2026 与 2025 两种标题格式。

    - 2026 标题『1.4-1.7\\n年货节（S-）』：等级在活动名括号内；
    - 2025 标题『活动-新年超值购\\nS-』或『活动-美力追新日\\nS-（银核预售）』：
      等级是独立行（`\\n` 后），`name` 里无括号等级。
    优先从 title 提取（title 总含等级信息），fallback 到 name 括号。
    """
    if title:
        m = re.search(r'[（(](S\+|S-|S|常规)[）)]', title)
        if m:
            return m.group(1)
        m = re.search(r'\n\s*(S\+|S-|S|常规)(?=[（(]|\s*$)', title)
        if m:
            return m.group(1)
    return _activity_grade(name)


def _activity_group_key(name: str) -> str:
    """归组键：含 × 复合名整体保留；否则去掉全部分隔符后，
    取前导活动代号（数字或 双11/双12 前缀：618 王楚钦→618、38现货→38、
    双11黄星官宣→双11），无则取整体。"""
    n = _normalize_activity_name(name)
    compact = re.sub(r'[\s\-—]+', '', n)
    if '×' in compact:
        return compact
    m = re.match(r'^(?:双\d+|\d+)', compact)
    return m.group(0) if m else (compact or name)


def _contiguous(prev_end: str, cur_start: str) -> bool:
    """两个同活动块是否时间连续：当前块起点 ≤ 上一块终点 + 1 天。"""
    try:
        a = datetime.date.fromisoformat(prev_end)
        b = datetime.date.fromisoformat(cur_start)
        return b <= a + datetime.timedelta(days=1)
    except Exception:  # noqa: BLE001
        return False


def _first_date(block: dict) -> str:
    d = block.get("dates") or []
    return d[0] if d else "9999-12-31"


def _group_waves(blocks: list) -> list:
    """把块按基础活动名归组，识别「同一活动的不同波段」。

    归组规则：
    - 名称归一化后相同的块先归一族；
    - 同族内若时间**不连续**（中间断档），按连续段拆成独立活动；
      （同一活动若在表里不是连续时间段，即使名字相同也视为独立活动）
    - 返回 activity_groups：每个连续段一个父活动，含其子波段（wave_count≥2 才算多波段）。
    - 给每个 block 打上 wave_group（归一化名）/ wave_no（段内序号）。
    """
    groups: dict = {}
    for i, b in enumerate(blocks):
        groups.setdefault(_activity_group_key(b["name"]), []).append(i)
    activity_groups = []
    wave_no_by_idx: dict = {}
    for k in sorted(groups, key=lambda k: _first_date(blocks[groups[k][0]])):
        idxs = sorted(groups[k], key=lambda i: _first_date(blocks[i]))
        # 按时间连续性切成子段
        sub_groups = []
        cur = []
        prev_end = None
        for i in idxs:
            d0 = blocks[i]["dates"][0] if blocks[i]["dates"] else None
            if cur and d0 and prev_end and not _contiguous(prev_end, d0):
                sub_groups.append(cur)
                cur = []
            cur.append(i)
            if d0:
                prev_end = blocks[i]["dates"][-1]
        if cur:
            sub_groups.append(cur)
        for sg in sub_groups:
            waves = [{
                "title": blocks[i]["title"],
                "name": blocks[i]["name"],
                "date_range": (f"{blocks[i]['dates'][0]}~{blocks[i]['dates'][-1]}"
                               if blocks[i]["dates"] else ""),
                "days": len(blocks[i]["dates"]),
            } for i in sg]
            for n, i in enumerate(sg, 1):
                wave_no_by_idx[i] = n
            activity_groups.append({"name": k, "wave_count": len(waves), "waves": waves})
    for i, b in enumerate(blocks):
        b["wave_group"] = _activity_group_key(b["name"])
        b["wave_no"] = wave_no_by_idx.get(i, 1)
    return activity_groups

def parse_page_blocks(ws):
    """解析「日常页面数据」或「活动页面数据」：按「时间」表头行切块。"""
    max_row = ws.max_row
    # 块边界：cell(r,1)=="时间" 的行是每个块的第一个表头行，块标题在其上一行
    time_rows = [r for r in range(1, max_row + 1)
                 if str(ws.cell(r, 1).value or "").strip() == "时间"]
    if not time_rows:
        return {"kind": "page_blocks", "fields": {}, "data": [], "sample": [],
                "anomalies": ["未识别到块结构：缺少「时间」表头行"]}

    blocks = []
    anomalies = []
    for idx, t in enumerate(time_rows):
        # 下一个块的标题行在 time_rows[idx+1]-1，当前块数据行应到其上一行（-2）
        t_end = time_rows[idx + 1] - 2 if idx + 1 < len(time_rows) else max_row
        block = _parse_one_block(ws, t, t_end)
        if block is None:
            anomalies.append(f"第 {idx + 1} 个块解析失败（行 {t} 起）")
        else:
            blocks.append(block)

    # 板块 / 子模块 清单
    blocks_all = set()
    subs_all = set()
    for b in blocks:
        for row in b["rows"]:
            if row["板块"]:
                blocks_all.add(row["板块"])
            if row["子模块"]:
                subs_all.add(row["子模块"])
    activity_groups = _group_waves(blocks)
    multi = [g for g in activity_groups if g["wave_count"] > 1]
    fields = {
        "block_count": len(blocks),
        "block_titles": [b["title"] for b in blocks],
        "板块": sorted(blocks_all),
        "子模块": sorted(subs_all),
        "layout": "A=板块 B=子模块 C/D=平均数据[点击率,支付转化] 每日期4列=新客[点击率,支付转化]·老客[点击率,支付转化]",
        "activity_groups": activity_groups,
        "multi_wave_count": len(multi),
    }
    if blocks:
        dates = blocks[0]["dates"]
        if not dates:
            anomalies.append("首个块未解析到日期列")
    return {"kind": "page_blocks", "fields": fields, "data": blocks,
            "sample": blocks[:2], "anomalies": anomalies}


def _extract_name(title: str) -> str:
    """块名提取，兼容多种标题：
    - 2026『1.4-1.7\\n年货节（S-）』『6.22\\n618返场（S）』『2.9 官宣后-2.11\\n欢聚日×年货节』
    - 2025『10.8-10.20 双11抢先购\\nS+』『活动-新年超值购\\nS-』
    """
    lines = [ln.strip() for ln in title.splitlines() if ln.strip()]
    if not lines:
        return title
    for ln in lines:
        s = ln.strip()
        if re.fullmatch(r'[（(]?S[+\-]?[）)]?', s):  # 等级行（S+/S-/S）跳过
            continue
        # 干净日期范围 + 名称（如 "10.8-10.20 双11抢先购"、"1.4-1.7"）
        m = re.match(r'^\d{1,2}\.\d{1,2}\s*[-~至]\s*\d{1,2}\.\d{1,2}\s*(.*)$', s)
        if m:
            rest = m.group(1).strip()
            if rest:
                return rest
            continue  # 纯日期范围行，名字在下一行
        # 以日期开头但不是干净范围（如 "2.9 官宣后-2.11"、"6.22"、"8.12-13"）→ 日期标签，跳过
        if re.match(r'^\d{1,2}\.\d{1,2}', s):
            continue
        return s
    return lines[0]


# 板块名规整：去形态后缀、明星专区归并、近似名归并
_STAR_ZONES = ("邓为", "王楚钦", "周柯宇")
_BLOCK_ALIAS = {
    "套装专区": "套装区",
    "新品区": "新品专区",
    "新品预热轮播": "新品专区",
    "爆品区": "爆款区",
    "爆品/主题专区": "爆款区",
    "银核新品专区": "银核专区",
    "银核新品派样": "银核专区",
}
_BLOCK_SUFFIX = ("单品", "套装", "TVC", "主推", "次推")


def _normalize_block_name(name: str) -> str:
    """把碎板块名归并成规整板块名：
    - 去括号后缀（单品/套装）与横杠后缀（-单品/-套装/-TVC/-主推/-次推）
    - 明星专区归并：邓为专区-* / 邓为主推专区 / 邓为玩法区 / 邓为连带同款 → 邓为专区
    - 近似名归并：套装专区→套装区、新品区→新品专区 等
    """
    s = (name or "").strip()
    if not s:
        return s
    s = re.sub(r'[（(](' + "|".join(_BLOCK_SUFFIX) + r')[)）]', '', s)
    for sfx in _BLOCK_SUFFIX:
        s = re.sub(r'[-]' + sfx + r'$', '', s)
    # 明星专区归并
    for star in _STAR_ZONES:
        if s.startswith(star):
            return f"{star}专区"
    return _BLOCK_ALIAS.get(s, s)


def _parse_one_block(ws, t, t_end):
    """解析从表头行 t（「时间」行）到 t_end 的一个块。兼容不同日期段宽度（2 列/4 列）。"""
    title = str(ws.cell(t - 1, 1).value or "").strip()
    name = _extract_name(title)

    h1, h2, h3 = t, t + 1, t + 2
    max_col = ws.max_column
    if str(ws.cell(h2, 1).value or "").strip() != "人群" \
            or str(ws.cell(h3, 1).value or "").strip() != "页面板块":
        return None

    # 人群分段列（人群行里的 新客/老客）
    seg_positions = []  # (col, "新客"/"老客")
    for c in range(3, max_col + 1):
        v = str(ws.cell(h2, c).value or "").strip()
        if v in ("新客", "老客"):
            seg_positions.append((c, v))

    def metric_at(c):
        m = str(ws.cell(h3, c).value or "").strip()
        return m or "点击率"

    # 平均数据列：C 起到首个日期段列之前（2026=点击率+支付转化 两列；2025=点击率 一列）
    avg_metrics = []  # (col, metric)
    if str(ws.cell(h1, 3).value or "").strip() == "平均数据":
        seg_start = seg_positions[0][0] if seg_positions else 3
        for c in range(3, min(seg_start, max_col + 1)):
            lbl = str(ws.cell(h3, c).value or "").strip()
            if not lbl:
                if c == 3:
                    lbl = "点击率"
                else:
                    break
            avg_metrics.append((c, lbl))

    # 日期段：每个「新客」列起一个段，段内按人群列切指标列
    date_blocks = []  # {iso, seg_ranges: {人群: (c0,c1)}}
    for i, (c_new, seg) in enumerate(seg_positions):
        if seg != "新客":
            continue
        j = i + 1
        while j < len(seg_positions) and seg_positions[j][1] != "新客":
            j += 1
        block_end = seg_positions[j][0] if j < len(seg_positions) else max_col + 1
        seg_ranges = {}
        for k in range(i, j):
            cc, lbl = seg_positions[k]
            nxt = seg_positions[k + 1][0] if k + 1 < j else block_end
            seg_ranges[lbl] = (cc, nxt)
        date_blocks.append({"iso": _date_to_iso(ws.cell(h1, c_new).value),
                            "seg_ranges": seg_ranges})
    if not date_blocks:  # 兜底：罕见布局，从 h1 直接扫日期，每段 2 列
        for c in range(3, max_col + 1):
            iso = _date_to_iso(ws.cell(h1, c).value)
            if iso:
                date_blocks.append({"iso": iso,
                                    "seg_ranges": {"新客": (c, c + 1), "老客": (c + 1, c + 2)}})

    # 数据行
    rows = []
    cur_block, cur_sub = None, None
    for r in range(t + 3, t_end + 1):
        a = _clean(ws.cell(r, 1).value)
        b = _clean(ws.cell(r, 2).value)
        if a is not None and not isinstance(a, (int, float)):
            cur_block = _normalize_block_name(a)
            cur_sub = None  # 进入新板块：子模块重置，等本行 B 再决定
        if b is not None and not isinstance(b, (int, float)):
            cur_sub = b
        avg = {}
        for c, m in avg_metrics:
            v = _clean(ws.cell(r, c).value)
            if v is not None:
                avg[m] = v
        daily = {}
        for db in date_blocks:
            iso = db["iso"]
            if iso is None:
                continue
            segs = {}
            for seg, (c0, c1) in db["seg_ranges"].items():
                mvals = {}
                for cc in range(c0, c1):
                    if cc > max_col:
                        break
                    v = _clean(ws.cell(r, cc).value)
                    if v is not None:
                        mvals[metric_at(cc)] = v
                segs[seg] = mvals
            daily[iso] = segs
        # 平均数据源给就用；没给则从每日新老客指标算术平均（忽略空值）
        avg_src = "源表" if avg else "程序计算"
        if not avg:
            buckets = {}
            for iso, segs in daily.items():
                for seg, mvals in segs.items():
                    for m, v in mvals.items():
                        if v is not None:
                            buckets.setdefault(m, []).append(v)
            avg = {m: sum(vs) / len(vs) for m, vs in buckets.items()}
        if not avg and cur_block is None and cur_sub is None and not daily:
            continue
        rows.append({"板块": cur_block, "子模块": cur_sub,
                     "平均数据": avg or None, "平均来源": avg_src, "每日": daily})

    dates = [db["iso"] for db in date_blocks if db["iso"]]
    return {"title": title, "name": name, "dates": dates,
            "row_count": len(rows), "rows": rows}


# ---------------- 入口 ----------------

def parse_file(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    known = {
        "整体数据": lambda ws: parse_overall(ws),
        "日常页面数据": lambda ws: parse_page_blocks(ws),
        "活动页面数据": lambda ws: parse_page_blocks(ws),
    }
    sheets = {}
    for name, fn in known.items():
        if name in wb.sheetnames:
            sheets[name] = fn(wb[name])
    unrecognized = [n for n in wb.sheetnames if n not in known]
    return {"sheets": sheets, "unrecognized_sheets": unrecognized}

